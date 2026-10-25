import json
import pathlib
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
from scipy import stats
from sklearn.linear_model import LinearRegression
from sklearn.metrics import r2_score
import openpyxl

XLSX    = pathlib.Path('../results/Tc_literature.xlsx')
RESULTS = pathlib.Path(RES)
EXCLUDE = {'NbTa67MoHfW33'}
ELEMENT_SET = ['Ti', 'Nb', 'Zr', 'Hf', 'Ta', 'Mo', 'W', 'Y', 'Re', 'Ru']   # set B; None → all 13 elements (incl. Sc, La, V)
INCLUDE_CAUTION = True     # False → only single-phase bcc rows
TC_MAIN = 'Tc_mu0.2'       # Tc used in the global and element-coloured plots
DEBYE_LABEL = 'θ_D mix'

KKR_KEYS = ['lambda', 'thetaDB_K', 'mixture_debye_temperature', 'Stoner_S', 'Stoner_E_sf_eff_meV']

# ── Load experimental Tc (Tc_literature.xlsx) + KKR results (results.json) ────
# xlsx sheet 'Tc_literature' (literature only, no KKR):
#   A=1 name, B–N = 2–14 composition (V = col 14 since 2026-10-23), S=19 Exp Tc, Z=26 First Author, AE=31 Status
# Section-header / empty rows (no Tc or no composition) are skipped.
# KKR results are matched by COMPOSITION (not by name): xlsx names and result
# directory names differ for some entries, and several literature rows share one
# composition (e.g. TaNbHfZr Krishnan / Wu2024) → they share one KKR result.

ELEMENTS   = ['Sc', 'Ti', 'Y', 'Zr', 'Nb', 'Mo', 'Ru', 'La', 'Hf', 'Ta', 'W', 'Re', 'V']   # xlsx columns B–N (V added 2026-10-23)
COL_NAME   = 1
COL_EL0    = 2
COL_TC_EXP = 19
COL_AUTHOR = 26
COL_STATUS = 31
COMP_TOL   = 2e-3   # max |Δx| for a composition match

def _val(cell):
    v = cell.value
    if v is None or v == 'NULL':
        return np.nan
    try:
        return float(v)
    except (TypeError, ValueError):
        return np.nan

def _num(v):
    return np.nan if v is None else v

# ── KKR results, keyed by directory name ──────────────────────────────────────
kkr = {}
for d in sorted(RESULTS.iterdir()):
    p = d / 'results.json'
    if not p.is_file() or p.stat().st_size == 0:
        continue
    j = json.loads(p.read_text())
    if not j:
        continue
    x = np.array([float(j.get(e) or 0.0) for e in ELEMENTS])
    kkr[d.name] = (x / x.sum(), j)

# ── literature rows ──────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb['Tc_literature']

rows, unmatched = [], []
for r in range(4, ws.max_row + 1):
    name   = ws.cell(r, COL_NAME).value
    tc_exp = _val(ws.cell(r, COL_TC_EXP))
    x      = np.nan_to_num([_val(ws.cell(r, COL_EL0 + i)) for i in range(len(ELEMENTS))])
    if not name or np.isnan(tc_exp) or x.sum() == 0:
        continue
    if name in EXCLUDE:
        continue
    x = x / x.sum()
    kdir = min(kkr, key=lambda k: np.abs(kkr[k][0] - x).max())
    if np.abs(kkr[kdir][0] - x).max() > COMP_TOL:
        unmatched.append(name)
        continue
    status  = ws.cell(r, COL_STATUS).value or ''
    caution = status.startswith('caution')
    if caution and not INCLUDE_CAUTION:
        continue
    j = kkr[kdir][1]
    rows.append(dict(
        name    = name,
        author  = ws.cell(r, COL_AUTHOR).value or '',
        status  = status,
        group   = 'caution (multiphase etc.)' if caution else 'single-phase bcc',
        Tc_exp  = tc_exp,
        kkr_dir = kdir,
        **dict(zip(ELEMENTS, x)),
        **{k: _num(j.get(k)) for k in KKR_KEYS},
    ))

df = pd.DataFrame(rows)
if ELEMENT_SET is not None:
    outside = [e for e in ELEMENTS if e not in ELEMENT_SET]
    n_all = len(df)
    df = df[(df[outside] == 0).all(axis=1)].reset_index(drop=True)
    print(f"ELEMENT_SET filter: kept {len(df)}/{n_all} rows (dropped alloys containing {outside})")
print(f"Loaded {len(df)} literature rows with KKR results "
      f"({(df.group != 'single-phase bcc').sum()} caution; excluded: {EXCLUDE}; "
      f"INCLUDE_CAUTION={INCLUDE_CAUTION})")
print(f"Unique KKR compositions used: {df.kkr_dir.nunique()}")
if unmatched:
    print(f"WARNING: no KKR result for: {unmatched}")

df[['name', 'author', 'group', 'Tc_exp', 'thetaDB_K', 'mixture_debye_temperature', 'lambda', 'kkr_dir']]

# ── Tc with mixture-rule θ_D ─────────────────────────────────────────────────
# λ_mix = λ_KKR·(θ_KKR/θ_mix)²  (η is θ_D-independent), then McMillan.
KB_MEV         = 0.08617
STONER_S_REF   = 1.15
STONER_MU_BASE = 0.185

def mcmillan(theta_D, lam, mu):
    denom = lam - mu * (1.0 + 0.62 * lam)
    if not np.isfinite(denom):
        return np.nan
    if denom <= 0.0:
        return 0.0
    return theta_D / 1.45 * np.exp(-1.04 * (1.0 + lam) / denom)

import sys
sys.path.insert(0, '../scripts')
from src.consts import STONER_I_EV, E_SF_MEV
RY_EV = 13.605693

def stoner_from_kkr(kdir):
    """(S, E_sf_mix) from per-component Ntot (states/Ry/atom/spin) in mcmillan_cutoff_results.csv.
    Same formula as process_kkr.compute_stoner_correction (C'-dependent γ disabled), but works for
    results.json written before the Stoner code."""
    rp = json.loads((RESULTS / kdir / 'run_params.json').read_text())
    c = pd.read_csv(RESULTS / kdir / 'finalscf' / 'mcmillan_cutoff_results.csv')
    c = c[(c.reduce_mode == 'mean') & (c.integral_mode == 'plain') & (c.norm_mode == 'none')]
    labels = rp['element_labels']
    conc = np.asarray(rp['concentrations'], float); conc /= conc.sum()
    IN, E_sf = 0.0, 0.0
    for _, r in c.iterrows():
        el, xi = labels[r.component - 1], conc[r.component - 1]
        IN   += xi * STONER_I_EV[el] * r.Ntot / RY_EV
        E_sf += xi * E_SF_MEV[el]
    return 1.0 / (1.0 - min(IN, 0.99)), E_sf

def stoner_mu_eff(S, E_sf_eff_meV, theta_D):
    # results.json from before the Stoner code (Jun–Jul runs) have no Stoner keys → NaN
    if not (np.isfinite(S) and np.isfinite(E_sf_eff_meV)):
        return np.nan
    gamma        = np.log(E_sf_eff_meV / (KB_MEV * theta_D))
    lsf          = S - 1.0
    mu_sf_full   = lsf / (1.0 + lsf * gamma)
    mu_sf_ref    = (STONER_S_REF - 1.0) / (1.0 + (STONER_S_REF - 1.0) * gamma)
    return STONER_MU_BASE + max(0.0, mu_sf_full - mu_sf_ref)

_se = [stoner_from_kkr(k) for k in df.kkr_dir]
_stored = df.Stoner_S.notna()
df['Stoner_S'] = [s for s, _ in _se]
df['Stoner_E_sf_eff_meV'] = [e for _, e in _se]
VEC_EL = {'Sc': 3, 'Y': 3, 'La': 3, 'Ti': 4, 'Zr': 4, 'Hf': 4, 'Nb': 5, 'Ta': 5, 'V': 5, 'Mo': 6, 'W': 6, 'Re': 7, 'Ru': 8}
df['VEC'] = sum(df[e] * v for e, v in VEC_EL.items())
df['theta_mix']  = df['mixture_debye_temperature']
df['lambda_mix'] = df['lambda'] * (df['thetaDB_K'] / df['theta_mix'])**2
for mu in (0.1, 0.2, 0.3):
    df[f'Tc_mu{mu}'] = [mcmillan(t, l, mu) for t, l in zip(df.theta_mix, df.lambda_mix)]
df['Tc_sf'] = [mcmillan(t, l, stoner_mu_eff(S, E, t))
               for t, l, S, E in zip(df.theta_mix, df.lambda_mix, df.Stoner_S, df.Stoner_E_sf_eff_meV)]

print(f"θ_D mix: {df.theta_mix.min():.0f}–{df.theta_mix.max():.0f} K, "
      f"λ_mix: {df.lambda_mix.min():.2f}–{df.lambda_mix.max():.2f}, "
      f"Tc_sf available for {df.Tc_sf.notna().sum()}/{len(df)} "
      f"(S recomputed from Ntot; {(~_stored).sum()} rows had no Stoner keys in results.json)")
df[['name', 'group', 'Tc_exp', 'theta_mix', 'lambda_mix', 'Tc_mu0.2', 'Tc_sf']].round(2)
